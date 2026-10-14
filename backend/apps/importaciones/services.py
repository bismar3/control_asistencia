import io
import openpyxl
from datetime import datetime, date, time
from django.utils import timezone


def procesar_excel_estudiantes(archivo_bytes):
    """
    Lee y valida el archivo Excel de estudiantes según formato:
    Columnas: Número de registro | Nombre completo | Correo electrónico
    """
    wb = openpyxl.load_workbook(io.BytesIO(archivo_bytes), data_only=True)
    sheet = wb.active

    estudiantes = []
    errores = []
    registros_vistos = set()

    # Identificar encabezados en la fila 1
    headers = [str(cell.value or '').strip().lower() for cell in sheet[1]]
    
    col_nro = -1
    col_nombre = -1
    col_email = -1

    for idx, h in enumerate(headers):
        if 'registro' in h or 'nro' in h or 'n°' in h or 'id' in h:
            col_nro = idx
        elif 'nombre' in h or 'estudiante' in h:
            col_nombre = idx
        elif 'correo' in h or 'email' in h:
            col_email = idx

    # Si no los detectó por nombre, usar posiciones por defecto (0, 1, 2)
    if col_nro == -1: col_nro = 0
    if col_nombre == -1: col_nombre = 1
    if col_email == -1: col_email = 2

    # Procesar filas a partir de la fila 2
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):  # Fila vacía
            continue

        raw_nro = row[col_nro] if col_nro < len(row) else None
        raw_nombre = row[col_nombre] if col_nombre < len(row) else None
        raw_email = row[col_email] if col_email < len(row) else None

        nro_registro = str(raw_nro).strip() if raw_nro is not None else ''
        nombre_completo = str(raw_nombre).strip() if raw_nombre is not None else ''
        email = str(raw_email).strip().lower() if raw_email is not None else ''

        fila_errores = []
        if not nro_registro:
            fila_errores.append("Número de registro es obligatorio.")
        elif nro_registro in registros_vistos:
            fila_errores.append(f"Número de registro duplicado en el archivo ({nro_registro}).")
        else:
            registros_vistos.add(nro_registro)

        if not nombre_completo:
            fila_errores.append("Nombre completo es obligatorio.")

        if not email or '@' not in email:
            fila_errores.append(f"Correo electrónico inválido o ausente ({email}).")

        if fila_errores:
            errores.append({
                'fila': row_idx,
                'nro_registro': nro_registro,
                'nombre_completo': nombre_completo,
                'email': email,
                'error': "; ".join(fila_errores)
            })
        else:
            # Separar nombre completo en nombre y apellidos
            partes = nombre_completo.split()
            if len(partes) == 1:
                nom = partes[0]
                pat = ""
                mat = ""
            elif len(partes) == 2:
                nom = partes[0]
                pat = partes[1]
                mat = ""
            elif len(partes) == 3:
                nom = partes[0]
                pat = partes[1]
                mat = partes[2]
            else:
                nom = " ".join(partes[:-2])
                pat = partes[-2]
                mat = partes[-1]

            estudiantes.append({
                'fila': row_idx,
                'nro_registro': nro_registro,
                'nombre_completo': nombre_completo,
                'nombre': nom,
                'apellido_paterno': pat,
                'apellido_materno': mat,
                'email': email
            })

    return estudiantes, errores


def procesar_excel_fechas(archivo_bytes):
    """
    Lee y valida el archivo Excel de fechas de clases según formato:
    Columnas: Fecha | Hora inicio | Hora fin
    """
    wb = openpyxl.load_workbook(io.BytesIO(archivo_bytes), data_only=True)
    sheet = wb.active

    clases = []
    errores = []
    fechas_vistas = set()

    headers = [str(cell.value or '').strip().lower() for cell in sheet[1]]
    col_fecha = -1
    col_ini = -1
    col_fin = -1

    for idx, h in enumerate(headers):
        if 'fecha' in h:
            col_fecha = idx
        elif 'inicio' in h or 'hora ini' in h:
            col_ini = idx
        elif 'fin' in h or 'hora fin' in h:
            col_fin = idx

    if col_fecha == -1: col_fecha = 0
    if col_ini == -1: col_ini = 1
    if col_fin == -1: col_fin = 2

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        raw_fecha = row[col_fecha] if col_fecha < len(row) else None
        raw_ini = row[col_ini] if col_ini < len(row) else None
        raw_fin = row[col_fin] if col_fin < len(row) else None

        parsed_fecha = None
        parsed_ini = None
        parsed_fin = None
        fila_errores = []

        # Parsear fecha
        if isinstance(raw_fecha, datetime):
            parsed_fecha = raw_fecha.date()
        elif isinstance(raw_fecha, date):
            parsed_fecha = raw_fecha
        elif isinstance(raw_fecha, str):
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%Y/%m/%d'):
                try:
                    parsed_fecha = datetime.strptime(raw_fecha.strip(), fmt).date()
                    break
                except ValueError:
                    pass
            if not parsed_fecha:
                fila_errores.append(f"Formato de fecha inválido ({raw_fecha}). Use DD/MM/YYYY.")
        else:
            fila_errores.append("Fecha es obligatoria.")

        if parsed_fecha:
            if parsed_fecha in fechas_vistas:
                fila_errores.append(f"Fecha duplicada en el archivo ({parsed_fecha.strftime('%d/%m/%Y')}).")
            else:
                fechas_vistas.add(parsed_fecha)

        # Parsear hora inicio
        if isinstance(raw_ini, time):
            parsed_ini = raw_ini
        elif isinstance(raw_ini, datetime):
            parsed_ini = raw_ini.time()
        elif isinstance(raw_ini, str):
            for fmt in ('%H:%M:%S', '%H:%M'):
                try:
                    parsed_ini = datetime.strptime(raw_ini.strip(), fmt).time()
                    break
                except ValueError:
                    pass

        # Parsear hora fin
        if isinstance(raw_fin, time):
            parsed_fin = raw_fin
        elif isinstance(raw_fin, datetime):
            parsed_fin = raw_fin.time()
        elif isinstance(raw_fin, str):
            for fmt in ('%H:%M:%S', '%H:%M'):
                try:
                    parsed_fin = datetime.strptime(raw_fin.strip(), fmt).time()
                    break
                except ValueError:
                    pass

        if fila_errores:
            errores.append({
                'fila': row_idx,
                'fecha': str(raw_fecha),
                'hora_inicio': str(raw_ini),
                'hora_fin': str(raw_fin),
                'error': "; ".join(fila_errores)
            })
        else:
            clases.append({
                'fila': row_idx,
                'fecha_programada': parsed_fecha,
                'hora_inicio': parsed_ini,
                'hora_fin': parsed_fin
            })

    return clases, errores
